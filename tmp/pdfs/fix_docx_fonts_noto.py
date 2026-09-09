from docx import Document
from docx.oxml.ns import qn
from pathlib import Path

src = Path('output/docx/マクロ経済学B_シラバス案.docx')
tmp = Path('tmp/pdfs/マクロ経済学B_シラバス案-fontfixed.docx')
doc = Document(src)
font_name = 'Noto Sans JP'


def set_rfonts(rpr):
    if rpr is None:
        return
    rfonts = rpr.rFonts
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rfonts.set(qn(f'w:{attr}'), font_name)


def set_run(run):
    run.font.name = font_name
    set_rfonts(run._element.get_or_add_rPr())


def set_paragraph(paragraph):
    for run in paragraph.runs:
        set_run(run)


def walk_table(table):
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                set_paragraph(p)
            for nested in cell.tables:
                walk_table(nested)

for style in doc.styles:
    if hasattr(style, 'font'):
        style.font.name = font_name
        set_rfonts(style._element.get_or_add_rPr())

for p in doc.paragraphs:
    set_paragraph(p)
for table in doc.tables:
    walk_table(table)
for section in doc.sections:
    for p in section.header.paragraphs + section.footer.paragraphs:
        set_paragraph(p)
    for table in section.header.tables + section.footer.tables:
        walk_table(table)

doc.save(tmp)
src.write_bytes(tmp.read_bytes())
print(src)
